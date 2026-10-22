"""Export the current app itinerary into a private workbook and printable handbook.

Original files are read only. Outputs must stay outside the public site and Git.
"""
from pathlib import Path
from copy import copy,deepcopy
from datetime import datetime
from collections import defaultdict
from html import escape as h
from http.server import ThreadingHTTPServer,SimpleHTTPRequestHandler
from threading import Thread
import argparse,json,re,hashlib,zipfile,math,textwrap
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.pagebreak import Break
from openpyxl.workbook.properties import CalcProperties
from PIL import Image,ImageDraw,ImageFont
from playwright.sync_api import sync_playwright
from travel_audit import review

ROOT=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser();parser.add_argument('--source',type=Path,required=True);parser.add_argument('--output',type=Path,required=True);parser.add_argument('--source-url',default='');parser.add_argument('--drive-folder-url',default='');parser.add_argument('--inventory',type=Path)
args=parser.parse_args();P=args.output.resolve()
if P.is_relative_to(ROOT):raise ValueError('私人文件必須輸出到Git倉庫外，不能放入網站目錄。')
P.mkdir(parents=True,exist_ok=True);(P/'assets').mkdir(exist_ok=True)
trip=json.loads((ROOT/'app/data/share.json').read_text());inventory=json.loads((args.inventory or P/'review/drive-ticket-inventory-1022.json').read_text());audit=review(trip,inventory)
original=openpyxl.load_workbook(args.source,keep_vba=True)
wb=openpyxl.load_workbook(args.source,keep_vba=True)
version=trip['version'];STAMP=version[5:7]+version[8:10];date=version[:10].replace('-','/');GREEN='173E38';PAPER='F6F5EF';ORANGE='FBEBDD';PURPLE='EFEDF8';GRAY='7C817A'
font='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
font_big=ImageFont.truetype(font,46,index=3);font_mid=ImageFont.truetype(font,28,index=3);font_small=ImageFont.truetype(font,23,index=3)

def link(url,label='開啟'):
    return {'url':url,'label':label} if url else ''
def href(cell):
    if cell.hyperlink:return cell.hyperlink.target
    m=re.search(r'HYPERLINK\("([^"]+)',str(cell.value or ''));return m.group(1) if m else ''
def detail(e):
    a=e.get('activity',{});parts=[a.get('description','')]
    for key,label in [('highlights','參觀重點'),('mustTry','必點／推薦'),('tips','注意')]:
        if a.get(key):parts.append(label+'：'+'；'.join(a[key]))
    if a.get('openingHours'):parts.append('營業：'+a['openingHours'])
    return '\n'.join(p for p in parts if p)
def zone_time(e):
    start=datetime.fromisoformat(e['start']);end=datetime.fromisoformat(e['end'])
    return e.get('displayTime') or start.strftime('%H:%M')+'–'+end.strftime('%H:%M')
def cards(day):
    stops=day.get('optionalStops',[]);result=[]
    for e in day['events']:
        result.extend(('optional',s) for s in stops if s.get('beforeEventId')==e['id'])
        result.append(('event',e))
        result.extend(('optional',s) for s in stops if s.get('afterEventId')==e['id'])
    result.extend(('optional',s) for s in stops if not s.get('beforeEventId') and not s.get('afterEventId'))
    return result
hotels=[]
for row in original['交通與住宿'].iter_rows(min_row=6,max_row=9):
    vals=[c.value for c in row]
    hotels.append({'city':vals[0],'name':vals[1],'dates':vals[2],'cost':vals[3],'address':vals[4],'reason':vals[5],'note':vals[6],'official':href(row[7]),'mapUrl':href(row[8])})
hotel_by_city={h['city']:h for h in hotels}
def hotel_name(e,day):
    names={'madrid':'馬德里','sevilla':'塞維亞','granada':'格拉納達','barcelona':'巴塞隆納'}
    if e.get('lodgingAction'):
        city=names.get(e.get('cityId'))
        return '住宿｜'+(hotel_by_city[city]['name']+'・' if city in hotel_by_city else '')+e['lodgingAction']
    return e.get('displayTitle',e['title'])
def color(e):
    if e.get('lodgingAction'):return PURPLE
    if e.get('category') in ['早餐','午餐','晚餐','甜點','條件甜點','餐飲散步']:return ORANGE
    return 'FFFFFF'

def table(name,title,headers,rows,widths=None,kinds=None):
    ws=wb[name] if name in wb.sheetnames else wb.create_sheet(name)
    for rg in list(ws.merged_cells.ranges):ws.unmerge_cells(str(rg))
    ws.delete_rows(1,ws.max_row);ws._images=[];ws.row_dimensions.clear();ws.freeze_panes='A6'
    n=len(headers);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n);ws['A1']=title
    ws['A1'].font=Font(name='Microsoft JhengHei',size=20,bold=True,color=GREEN);ws.row_dimensions[1].height=34
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=n);ws['A2']=date+' 同步｜跨國航班依起降當地時間；「目標」非已購；淡字備選不安排時間。'
    ws['A2'].font=Font(name='Microsoft JhengHei',size=10,color=GRAY);ws.row_dimensions[2].height=25
    for c,val in enumerate(headers,1):
        cell=ws.cell(5,c,val);cell.fill=PatternFill('solid',fgColor=GREEN);cell.font=Font(name='Microsoft JhengHei',size=11,bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[5].height=30
    for i,row in enumerate(rows,6):
        kind=(kinds or {}).get(i-6,'');fill=GRAY if kind=='optional' else GREEN
        length=0
        for c,val in enumerate(row,1):
            if isinstance(val,datetime):val=val.replace(tzinfo=None)
            cell=ws.cell(i,c)
            if isinstance(val,dict) and 'url' in val:
                cell.value=val['label'];cell.hyperlink=val['url'];cell.font=Font(name='Microsoft JhengHei',size=10,color='365E91',underline='single')
            else:
                cell.value=val;cell.font=Font(name='Microsoft JhengHei',size=10,color=GRAY if kind=='optional' else '243D36')
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            cell.fill=PatternFill('solid',fgColor='F3F2ED' if kind=='optional' else kind if kind in [ORANGE,PURPLE] else ('F8F8F2' if i%2 else 'FFFFFF'))
            cell.border=Border(bottom=Side(style='hair',color='DADFD5'))
            if isinstance(val,datetime):cell.number_format='m/d hh:mm' if c in [3,4] and name=='詳細時間流' else 'yyyy/m/d'
            length=max(length,len(str(val or '')))
        ws.row_dimensions[i].height=min(126,max(32,22+12*math.ceil(length/75)))
    for c in range(1,n+1):ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width=(widths or {}).get(c,24)
    ws.auto_filter.ref=f'A5:{openpyxl.utils.get_column_letter(n)}{max(6,5+len(rows))}'
    ws.print_title_rows='1:5';ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_options.horizontalCentered=True;ws.sheet_view.showGridLines=False
    return ws

# Refresh maps as clearly labelled route diagrams, with navigation links kept in tables.
def wrap(draw,text,font,max_width):
    lines=[];line=''
    for char in text:
        if draw.textlength(line+char,font=font)>max_width and line:lines.append(line);line=char
        else:line+=char
    if line:lines.append(line)
    return lines
posters={}
for day in trip['days']:
    points=[];seen=set()
    for kind,obj in cards(day):
        if kind=='optional':points.append(('有空',obj.get('displayTitle',obj['title']),True));continue
        for e in obj.get('mealOptions') or [obj]:
            if not e.get('venue') and not e.get('isFlight'):continue
            key=e.get('venue',{}).get('nameEn') or e['title']
            if key in seen:continue
            seen.add(key);title=e.get('title',obj['title']) if obj.get('mealOptions') else obj.get('displayTitle',obj['title'])
            points.append((e.get('displayTime') or obj['start'][11:16],title,False))
    image=Image.new('RGB',(1400,820),'#F6F5EF');draw=ImageDraw.Draw(image)
    draw.text((55,30),day['date']+'  '+day['city'],font=font_big,fill='#173E38')
    draw.text((55,100),'最新動線示意｜非等比例街道圖，實際移動請開啟導航',font=font_small,fill='#6E7D73')
    cols=3;rows=max(1,math.ceil(len(points)/cols));boxh=min(155,570/rows);boxw=415
    for i,(time,title,optional) in enumerate(points):
        col=i%cols;row=i//cols;x0=55+col*445;y0=155+row*(boxh+8)
        draw.rounded_rectangle((x0,y0,x0+boxw,y0+boxh),radius=16,fill='#ECEBE6' if optional else '#E3EBDD')
        draw.text((x0+15,y0+8),str(i+1).zfill(2)+'  '+time[:26],font=font_small,fill='#788277' if optional else '#4F6A51')
        lines=wrap(draw,title,font_mid,boxw-30)
        for j,line in enumerate(lines[:max(1,int((boxh-42)/33))]):draw.text((x0+15,y0+41+j*33),line,font=font_mid,fill='#7B7E75' if optional else '#173E38')
    draw.text((55,775),version+'｜目標時段以最後票面為準；有空地點不占固定時段',font=font_small,fill='#7B8278')
    file=P/'assets'/f'{day["date"]}-route.png';image.save(file,optimize=True);posters[day['date']]=file

summary=[];timeline=[];timeline_kinds={};movement=[]
for index,day in enumerate(trip['days'],1):
    main=[hotel_name(e,day) for e in day['events'] if e.get('venue') or e.get('isFlight')]
    sections=[]
    for low,high in [(0,12),(12,18),(18,24)]:
        sections.append('；'.join(e.get('displayTitle',e['title']) for e in day['events'] if low<=int(e['start'][11:13])<high and (e.get('venue') or e.get('isFlight'))))
    city=next((city for city in hotel_by_city if city in day['city']),None)
    summary.append([datetime.fromisoformat(day['date']),['一','二','三','四','五','六','日'][datetime.fromisoformat(day['date']).weekday()],day['city'],day['title'],*sections,hotel_by_city[city]['name'] if city else '航程／返抵日','以已購票／已訂位優先',link('https://huang-8x1z.vercel.app/#plan','App行程')])
    for kind,obj in cards(day):
        optional=kind=='optional';start=None if optional else datetime.fromisoformat(obj['start']);end=None if optional else datetime.fromisoformat(obj['end'])
        notes=(obj.get('description','')+'\n'+obj.get('note','')+'\n'+'；'.join(obj.get('mustTry',[]))) if optional else detail(obj)
        if not optional and obj.get('mealOptions'):
            notes+='\n'+ '\n'.join(o['title']+'｜'+o['displayTime'] for o in obj['mealOptions'])
        if not optional and obj.get('isFlight'):notes+='\n起飛時區：'+obj.get('timezone','Europe/Madrid')+'；抵達時區：'+obj.get('endTimezone','Europe/Madrid')
        recs=([obj['reservation']] if obj.get('reservation') else [])+[o['reservation'] for o in obj.get('mealOptions',[]) if o.get('reservation')]+([obj['ticketPlan']] if obj.get('ticketPlan') else [])
        timeline_kinds[len(timeline)]='optional' if optional else color(obj)
        timeline.append([datetime.fromisoformat(day['date']),f'D{index}',start,end,None if optional else round((end-start).total_seconds()/60),'彈性備選' if optional else '住宿' if obj.get('lodgingAction') else obj['category'],obj.get('displayTitle',obj['title']) if optional else hotel_name(obj,day),'依導航／司機確認' if not optional and obj['category']=='交通' else '',notes.strip(),'、'.join(r['id'] for r in recs),link(obj.get('mapUrl'),'導航'),obj.get('venue',{}).get('nameEn',''),obj.get('status',''),'','','',obj.get('id','')])
        if not optional and obj['category'] in ['交通','高鐵','航班']:
            movement.append([f'L{len(movement)+1:03}',f'D{index}',day['date'],zone_time(obj),'依上一站',hotel_name(obj,day),obj.get('title',''),'',link(obj.get('ticketPlan',{}).get('url'),'購票／官方'),link(obj.get('mapUrl'),'導航'),detail(obj)])
table('行程總覽','西班牙口袋旅程｜13個日期、11天西班牙行程',['日期','星期','城市','主題','上午','下午','晚上','住宿','優先順序','連結'],summary,{4:35,5:50,6:50,7:45,8:43})
wb['行程總覽']['A3']='費用請看預算頁；原版預算尚未依本次新增博物館與餐廳完整重估。';wb['行程總覽']['A3'].font=Font(name='Microsoft JhengHei',color='946129',size=10)
table('詳細時間流','最新版完整時間軸｜淡字地點不鎖時間',['日期','天次','開始（當地）','結束（當地）','分鐘','類別','活動／目的地','交通','活動資訊／必點／注意','訂位／票券ID','導航','店家／景點','狀態','交通管道','預買／預約','官方','App活動ID'],timeline,{1:14,2:8,3:18,4:18,5:9,6:14,7:64,8:24,9:100,10:24,11:12,12:38,13:36},timeline_kinds)
# Accommodations retain original quoted prices and addresses; update obsolete handling notes.
for hotel in hotels:
    if hotel['city']=='馬德里':hotel['reason']='12/24寄物後逛市場；12/26飯店已回覆可寄放28吋大箱，先去王宮，下午回來領箱，不使用Locker Prado。';hotel['note']='12/26飯店已確認可寄放行李（使用者提供）。12/24早到寄物／房間交付及假日接待仍依住宿回覆。'
    if hotel['city']=='塞維亞':hotel['note']='12/26列車21:39抵達，先入住，22:30 El Rinconcillo只保留10分鐘。接車、晚到入住與門碼請事先確認；12/28清晨退房，06:45接車、08:00巴士。'
    if hotel['city']=='格拉納達':hotel['note']='12/28 11:35寄物、18:40入住；12/29 07:45大車往GRX，10:55航班。早餐與四箱接送需準備。'
    if hotel['city']=='巴塞隆納':hotel['note']='12/29下飛機後直接到住宿放行李，是否可提早入住依住宿回覆；1/3 06:15清晨接車，約07:00到BCN T1。'
table('交通與住宿','住宿與行李處理｜原房費報價保留，不推定已付款',['城市','住宿名稱','入住→退房','原報價','地址','行程銜接','入住／寄物／接車提醒','官方','導航'],[[q['city'],q['name'],q['dates'],q['cost'],q['address'],q['reason'],q['note'],link(q['official']),link(q['mapUrl'])] for q in hotels],{2:49,4:33,5:63,6:70,7:90})
mealrows=[];foodguide=[];choices=[]
for i,r in enumerate(audit['meal_records'],1):
    a=r['activity'];must='；'.join(a.get('mustTry',[]));notes='；'.join(a.get('tips',[]));m=link(r['mapUrl'],'導航')
    mealrows.append([r['date']+' '+r['time'],next(d['city'] for d in trip['days'] if d['date']==r['date']),r['title'],must,'依導航核對分店','價格待確認',r['status'],r['kind']+'；'+notes,'App '+version,m])
    foodguide.append([f'FOOD{i:02}',r['date']+' '+r['time'],r['title'],must,'價格待確認',r['status'],a.get('description','')+'；'+notes,'依導航核對分店','',m,m])
    choices.append([r['sourceId'],r['date'],r['time'],r['kind'],r['name'],a.get('description',''),r['status'],must,'','','','','依導航核對分店',notes,'已購景點與訂位優先',m,m,m])
table('餐廳與美食','目前餐飲安排｜主選、擇一與彈性分開',['日期／時段','城市','店家','推薦／必點','分店','價格','狀態','注意事項','來源','導航'],mealrows,{1:29,3:58,4:75,7:40,8:100})
table('餐飲詳解','目前餐飲與店家資訊｜未核實菜單保留說明',['ID','安排','店家','推薦／必點','價格','訂位狀態','活動資訊與注意','分店','官方','店家位置','導航'],foodguide,{2:30,3:59,4:72,7:110})
table('餐次選擇','餐次選擇｜兩間已訂午餐仍只算一餐',['餐次ID','日期','時段','類型','店家','介紹','狀態','推薦','EUR低','EUR高','TWD低','TWD高','地址／分店','注意事項','優先順序','店家','聯絡','導航'],choices,{1:32,4:20,5:49,6:68,8:70,14:90})
# Ticket ledger: retain paid figures in original rows and supplement current Drive links.
tickets=[]
for row in original['票券核對'].iter_rows(min_row=6,max_row=original['票券核對'].max_row):
    vals=[c.value for c in row]
    if vals[0]:tickets.append(dict(zip(['id','date','title','time','report','people','paid','eur','status','note'],vals[:10]),url=href(row[10])))
byid={t['id']:t for t in tickets}
byid['K08'].update(status='已購；10/22重讀原票',note='08:00 Sevilla Plaza de Armas→11:00 Granada Bus Station；原票四人，兩端大車接駁另確認。',url=audit['transport'][2]['url'])
byid['K09'].update(time='10:55–12:20',status='已購，先前原行程單已核對',note='每人25kg托運及40×20×30座下物品；四人總收據NT$35,691含TripCoins折抵，沿用既有同行包原件。',paid=35691)
byid['K10'].update(status='晚餐未安排',time='視主人聯絡；否則住宿自理',note='使用者尚未聯絡到主人，不視為已訂位。')
inv=json.loads((P/'review/drive-ticket-inventory-1022.json').read_text())
def drivefind(text):
    item=next((f for f in inv if text in f['path'] and f['mime']=='application/pdf'),None)
    return 'https://drive.google.com/file/d/'+item['id']+'/view' if item else ''
new=[('K11','2026-12-29','巴特略之家鉑金票','16:30','16:10','已購','四人鉑金；加購依每張票不同，不推定全員含屋頂甜點／酒。',drivefind('加購服務')),('K12','2026-12-30','聖保祿醫院','10:00','09:40','已購','四人自助參觀與英語語音，依既有票面。',drivefind('聖保祿醫院_票券')),('K13','2026-12-31','奎爾公園','09:30','09:00','已購','固定09:30優先，與12:00文森之家銜接。',drivefind('奎爾公園票券')),('K14','2026-12-31','文森之家','12:00','11:45','已購','四人票面固定時段。',drivefind('文森之家票券')),('K15','2027-01-01','米拉之家','11:30','11:00','已購','元旦票面固定，Mauri有空才去。',drivefind('米拉之家 Tickets')),('K16','2027-01-02','奎爾紡織村地下教堂','10:00目標','09:40目標','日期錯誤待改訂','10/22重讀原票仍02/11/2026 10:00；目標2027/1/2，尚無正確新票；不推定含FGC。',drivefind('奎爾紡織村地下教堂 Tickets')),('K17','2026-12-26','Renfe AVE 02190 馬德里→塞維亞','19:00–21:39','18:15目標','已購；10/22重讀原票','車廂2，座位6B/6C/7B/7C；票面€121.20，先前提供€100.75，差€20.45待對帳，預算不擅改。',audit['transport'][1]['url']),('F01','2026-12-23','國際機票・中國東方航空','12/23 18:40出發；1/4 14:00返抵','各機場當地時間','已開票','MU5006／MU709／MU250／MU5007；沿用先前提供的原行程單。','')]
for id,dt,title,time,report,status,note,url in new:byid[id]={'id':id,'date':dt,'title':title,'time':time,'report':report,'people':4,'paid':'','eur':'','status':status,'note':note,'url':url}
bookings=[]
for day in trip['days']:
 for e in day['events']:
    bookings.extend(([e['reservation']] if e.get('reservation') else [])+[o['reservation'] for o in e.get('mealOptions',[]) if o.get('reservation')]+([e['ticketPlan']] if e.get('ticketPlan') else []))
# User-confirmed restaurant bookings that precede the app's structured records.
for id,dt,name,time in [('RLOS','2026-12-25','Los Galayos','19:00'),('RBOT','2026-12-26','Botín','13:00'),('RELR','2026-12-26','El Rinconcillo','22:30'),('RPOST','2026-12-27','Bar Postiguillo Tapas','13:00')]:
 bookings.append({'id':id,'date':dt,'title':name,'time':time,'report':'提早到店','status':'已訂位・使用者確認','note':'依使用者先前確認；確認信尚未提供。'+('只保留10分鐘至22:40。' if id=='RELR' else ''),'url':''})
for b in bookings:byid[b['id']]={'id':b['id'],'date':b['date'],'title':b['title'],'time':b['time'],'report':b['report'],'people':'依訂位／票面','paid':'','eur':'','status':b['status'],'note':b['note'],'url':b.get('url','')}
prado=audit['attractions'][0];byid['P01']={'id':'P01','date':'2026-12-26','title':'普拉多美術館','time':'15:45–17:00目標','report':'15:25目標','people':4,'paid':'','eur':'','status':'尚未購票','note':prado['action'],'url':prado['url']}
ledger=sorted(byid.values(),key=lambda t:(str(t['date']),t['id']))
table('票券核對','票券與訂位核對｜沒有原票不等於未付款',['ID','使用日期','項目','入場／交通','報到／到站','人數','已付團TWD','票面／參考EUR','狀態','注意事項','原票／購票'],[[t[k] for k in ['id','date','title','time','report','people','paid','eur','status','note']]+[link(t['url'])] for t in ledger],{3:47,4:40,9:43,10:115})
all_pending=audit['attractions']+audit['restaurants_pending']+[t for t in audit['transport'] if not t['status'].startswith(('已購','已開票'))]
preorders=[[f'P{i:02}', '優先' if i<5 else '確認',q['date'],'',q['item'],'依票面／店家條款','依行程與開放時間',q['status'],'沒有憑證時標記待確認，不斷言未訂',q['action'],link(q['url'])] for i,q in enumerate(all_pending,1)]
table('預訂清單','待購票、改訂與接送確認',['ID','優先','使用日','時段','項目','先決條件','何時處理','狀態','依據','具體操作','連結'],preorders,{3:30,5:51,8:44,9:48,10:115})
table('交通管道','交通票券與接送核對',['ID','服務','具體操作','行李／銜接','備案','連結','來源','完成狀態'],[[f'T{i:02}',t['item'],t['action'],'4人4大箱接駁需確認容量' if '大車' in t['item'] else '依票面與實際班次','市內車可現場叫車；清晨／緊接訂位段先確認',link(t['url']),audit['scope'],t['status']] for i,t in enumerate(audit['transport'],1)],{2:57,3:110,4:52,7:80,8:48})
table('交通預訂','最新版所有移動｜導航不是已預約憑證',['移動ID','天次','日期','當地時段','起點','終點','移動標題','管道','官方／購票','導航','注意'],movement,{4:36,6:70,7:65,11:105})
# Current attraction guide; repeated areas are grouped by their first appearance.
guide=[];seen=set()
for day in trip['days']:
 for e in day['events']:
    if not e.get('venue') or e['category'] in ['早餐','午餐','晚餐','甜點','條件甜點','採買']:continue
    key=e['venue']['nameEn']
    if key in seen:continue
    seen.add(key);a=e.get('activity',{});guide.append([f'A{len(guide)+1:02}',day['date']+' '+zone_time(e),e['venue']['nameZh']+'（'+key+'）',a.get('description',''),'；'.join(a.get('highlights',[])),'；'.join(a.get('tips',[])),a.get('photoTip',''),e.get('status',''),'費用依票種',key,link(e.get('ticketPlan',{}).get('url')),link(e.get('ticketPlan',{}).get('url')),link(e.get('mapUrl'))])
table('景點指南','目前景點指南｜含畢卡索博物館',['ID','安排','景點','介紹','參觀重點','提醒','拍照','票況','費用','區域','官方','購票','導航'],guide,{2:35,3:65,4:100,5:75,6:115,7:65})
table('行前與備案','行前與彈性處理',['日期／時段','分類','具體操作','相關','來源','狀態'],[[q['date'],'待訂／確認',q['action'],q['item'],'10/22核對',q['status']] for q in all_pending]+[['1/1','午餐擇一','Margherita13:45／El Furacu14:00都已訂擇一；決定後依條款處理未用訂位。','元旦餐飲','使用者確認','待決定'],['1/2','彈性午餐','奎爾宮未購、放票後可提早；午餐可跳過或臨時簡餐，市場海鮮不是必須吃。','奎爾宮→畢卡索→19:00 Sagardi','使用者10/22','彈性']],{1:31,3:115,4:64,6:45})
sources=[['U1022','使用者10/22調整','1/2奎爾宮可提前／午餐彈性、Re-Cuir、畢卡索、19:00 Sagardi已訂位',date,''],['APP','已部署App '+version,'本次Excel與PDF同一份行程資料',date,link('https://huang-8x1z.vercel.app/')],['X1013','原Excel','保留原始檔，最新版由本程式同步重編；原檔未覆寫',date,link(args.source_url)],['DRV','雲端票券目錄','10/22重新列出；未發現奎爾宮或畢卡索票；未讀取證件資料夾',date,link(args.drive_folder_url)],['PALAU','奎爾宮官方','冬季10:00–17:30、最後入場16:30；未證實1/2已放票',date,link('https://inici.palauguell.cat/en/')],['PICASSO','畢卡索博物館官方','目前冬季週二至週日10:00–19:00；實際票面優先',date,link('https://museupicassobcn.cat/en/plan-your-visit')]]
table('來源與訊息','本次更新來源與核對範圍',['來源ID','來源','核對內容','查核日期','連結'],sources,{2:42,3:120,4:17})
table('餐飲檢視','重複與類型檢視｜候選不重複計算',['主題','觀察','建議'],[[f['topic'],f['result'],f['advice']] for f in audit['findings']],{1:27,2:115,3:115})
table('待訂與確認','目前待訂與交通核對',['分類','項目','日期／目標','狀態','下一步','連結'],[[cat,q['item'],q['date'],q['status'],q['action'],link(q['url'])] for cat,group in [('景點',audit['attractions']),('餐廳',audit['restaurants_pending']),('交通',audit['transport'])] for q in group],{1:12,2:59,3:38,4:50,5:115})
jan2=next(d for d in trip['days'] if d['date']=='2027-01-02')
table('1月2日動線','1/2可調整動線｜13:00與15:30均未購目標',['時段','類型','活動','狀態','注意','導航'],[['無固定時間' if k=='optional' else zone_time(o),'彈性' if k=='optional' else o['category'],o.get('displayTitle',o['title']),o.get('status',''),o.get('note','') if k=='optional' else detail(o),link(o.get('mapUrl'))] for k,o in cards(jan2)],{1:35,2:13,3:85,4:44,5:120})
ws=wb['每日地圖']
for rg in list(ws.merged_cells.ranges):ws.unmerge_cells(str(rg))
ws.delete_rows(1,ws.max_row);ws._images=[];ws.row_dimensions.clear();ws.sheet_view.showGridLines=False
for idx,day in enumerate(trip['days']):
 row=idx*45+1;ws.cell(row,1,day['date']+' '+day['city']+'｜'+day['title']).font=Font(name='Microsoft JhengHei',size=17,bold=True,color=GREEN)
 ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=14)
 ws.cell(row+1,1,'動線示意非等比例地圖；最新時間與導航請見詳細時間流，淡字地點有空再去。')
 ws.merge_cells(start_row=row+1,start_column=1,end_row=row+1,end_column=14)
 img=XLImage(posters[day['date']]);img.width=1050;img.height=615;ws.add_image(img,f'A{row+3}')
 if idx:ws.row_breaks.append(Break(id=row-1))
for c in range(1,15):ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width=12
ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
# Preserve original budget formulas and inputs as reference; do not pretend the new restaurants are priced.
wb['預算']['A1']='原版預算參考｜尚未依本次新票與餐廳完整重估'
wb['預算']['A2']='此頁保留原1013預算輸入與公式，部分交通與票況已更新於票券核對／待訂與確認；新增畢卡索及新餐廳價格未包含，不視為最新總額。'
wb['預算']['A2'].font=Font(name='Microsoft JhengHei',size=11,bold=True,color='A55D28')
wb.calculation=CalcProperties(fullCalcOnLoad=True,forceFullCalc=True)
workbook=P/f'Spain_2026_Travel_Timeline_{STAMP}.xlsm';wb.save(workbook)

# Human review and ticket link index, kept private alongside the downloadable workbook.
md=['# 西班牙行程與預訂檢視｜2026/10/22','',audit['scope'],'','## 餐廳重複與類型','']
for f in audit['findings']:md += ['### '+f['topic'],'',f['result']+' '+f['advice'],'']
for label,group in [('交通票券與接送',audit['transport']),('景點票券',audit['attractions']),('仍待預約餐廳',audit['restaurants_pending'])]:
 md+=['## '+label,'']
 for q in group:md+=['- **'+q['item']+'｜'+q['date']+'｜'+q['status']+'**：'+q['action']+' [連結]('+q['url']+')','']
md+=['## 文件與費用說明','','原Excel未覆寫，新版同步目前App。原1013預算頁保留為參考，尚未完整重估新增博物館與新餐廳價差。Renfe票面與先前提供金額分列，沒有擅改預算。','']
(P/f'Spain_2026_Review_{STAMP}.md').write_text('\n'.join(md))
(P/f'Spain_2026_Review_{STAMP}.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
linkhtml='<meta charset="utf-8"><title>西班牙票券連結</title><h1>票券與訂位連結｜2026/10/22</h1><p>日期錯誤票券不能視為有效；餐廳確認信未提供者以使用者確認標示。</p><table border="1" cellpadding="8"><tr><th>日期</th><th>項目</th><th>時間</th><th>狀態</th><th>注意／連結</th></tr>'+''.join('<tr>'+''.join('<td>'+h(str(t[k] or ''))+'</td>' for k in ['date','title','time','status'])+'<td>'+h(t['note'])+(f'<br><a href="{h(t["url"],quote=True)}">原票／官方</a>' if t['url'] else '')+'</td></tr>' for t in ledger)+'</table>'
(P/f'Spain_2026_Ticket_Links_{STAMP}.html').write_text(linkhtml)

def a(url,label):return f'<a href="{h(url,quote=True)}">{h(label)}</a>' if url else ''
def li(items):return '<ul>'+''.join('<li>'+h(str(i))+'</li>' for i in items if i)+'</ul>' if items else ''
def info(e,logistics=False):
 act=e.get('activity',{});out=''
 if act.get('description') and not logistics:out+='<p>'+h(act['description'])+'</p>'
 for key,label in [('highlights','參觀重點'),('mustTry','必點／推薦'),('tips','注意')]:
  if act.get(key):out+='<p class="small"><b>'+label+'</b></p>'+li(act[key])
 return out
html=['''<!doctype html><html lang="zh-Hant"><meta charset="utf-8"><title>西班牙旅行手冊｜2026/10/22</title><style>
@page{size:A4;margin:16mm 15mm 18mm}*{box-sizing:border-box}body{font-family:"Noto Sans CJK TC","Noto Sans CJK",sans-serif;font-size:10px;line-height:1.65;color:#243d36;margin:0}h1{font-size:24px;line-height:1.45;color:#173e38;margin:8px 0 12px}h2{font-size:16px;line-height:1.5;color:#173e38;margin:16px 0 9px;break-after:avoid}h3{font-size:12px;margin:8px 0 5px;break-after:avoid}p{margin:5px 0}a{color:#386e67;text-decoration:underline;overflow-wrap:anywhere}ul{padding-left:18px;margin:5px 0}li{margin:3px 0}.cover{padding:40px 18px 35px;min-height:880px}.eyebrow{font-size:10px;letter-spacing:2px;color:#6c8176}.cover h1{font-size:38px;margin-top:35px}.notice{background:#f4ecd9;padding:13px;border-left:4px solid #ba8953;margin:12px 0}.section{break-before:page}.day{break-before:page}.route{width:100%;margin:10px 0 15px}.event{display:flex;gap:12px;padding:10px 10px;border-bottom:1px solid #dce2d6;break-inside:avoid}.event .time{width:90px;flex-shrink:0;font-size:11px;font-weight:600}.body{flex:1;min-width:0}.event h3{margin-top:0}.dining{background:#fcf0e4}.lodging{background:#efedf8}.optional{margin:9px 0 9px 102px;background:#f0efeb;color:#777d74;padding:10px 13px;border-left:3px solid #bbbdb5;break-inside:avoid}.optional h3{color:#777d74}.status{display:inline-block;padding:2px 6px;background:#e6ecdf;border-radius:4px;margin:4px 0;font-size:9px}.warning{background:#f4e1c0;color:#946129}.small{font-size:9px;color:#6a796d}.box{padding:13px;border:1px solid #dce2d6;margin:9px 0;break-inside:avoid}.box h3{font-size:13px}.meal-option{border-left:3px solid #b88052;background:#fff8ef;margin:8px 0;padding:8px 11px}table{border-collapse:collapse;width:100%;margin:12px 0}td,th{border:1px solid #dde3d8;text-align:left;vertical-align:top;padding:7px;overflow-wrap:anywhere}th{background:#e5ebdf}tr{break-inside:avoid}.key{font-weight:600;color:#173e38}.flow{font-size:13px;line-height:1.9}.muted{color:#7a857b}
</style><body>''']
html+=['<section class="cover"><p class="eyebrow">SPAIN · CHRISTMAS & NEW YEAR · FOUR TRAVELLERS</p><h1>西班牙<br>口袋旅行手冊</h1><h2>2026/12/23 出發 — 2027/1/4 返抵台灣</h2><p>西班牙11天 · 馬德里／塞維亞／格拉納達／巴塞隆納</p><p>更新：'+date+' · '+h(version)+'</p><div class="notice"><b>本次重點</b><p>1/2 奎爾宮13:00與畢卡索15:30皆為未購彈性目標；午餐可略過或臨時簡餐，途中可逛Re-Cuir，19:00 Sagardi已訂位。</p><p>原地下教堂PDF仍為02/11/2026，尚待改成2027/1/2；不能當作本次有效票。</p></div><p>'+a('https://huang-8x1z.vercel.app/','開啟手機App')+' · '+a(args.drive_folder_url,'開啟雲端票券資料夾')+'</p><p class="small">本手冊依最新版App、原Excel及本次票券核對重編。原檔保留未覆寫；住宿與原票連結僅放在本私人文件，不公開到網站。</p><h2>使用方式</h2><p>已購票／已訂位優先，未購時段都是目標。淡字地點不指定時間，有空再去。導航及雲端原票需網路；同行包內的票券仍可按既有App方式離線開啟。</p><p>跨國航段依各機場當地時間；西班牙冬季為UTC+1。台灣出發12/23 18:40，返抵台灣1/4 14:00。</p></section>']
html+=['<section class="section"><h1>餐飲與待辦快覽</h1>']
for f in audit['findings']:html+=['<article class="box"><h3>'+h(f['topic'])+'</h3><p>'+h(f['result'])+'</p><p><b>建議：</b>'+h(f['advice'])+'</p></article>']
html+=['</section><section class="section"><h1>尚待購票、改訂與交通確認</h1><p class="small">'+h(audit['scope'])+'</p>']
for label,group in [('景點票券',audit['attractions']),('餐廳',audit['restaurants_pending']),('交通票券／接送',audit['transport'])]:
 html+=['<h2>'+label+'</h2>']
 for q in group:html+=['<article class="box"><h3>'+h(q['item'])+'</h3><p><b>'+h(q['date'])+'｜'+h(q['status'])+'</b></p><p>'+h(q['action'])+'</p>'+a(q['url'],'官方／原票／安排入口')+'</article>']
html+=['</section><section class="section"><h1>住宿與行李</h1>']
for q in hotels:html+=['<article class="box"><p class="eyebrow">'+h(q['city']+' · '+q['dates'])+'</p><h2>'+h(q['name'])+'</h2><p>'+h(q['address'])+'</p><p>'+h(q['note'])+'</p><p class="small">原報價：'+h(str(q['cost']))+'；未推定付款狀態。</p>'+a(q['mapUrl'],'住宿導航')+' · '+a(q['official'],'住宿官網')+'</article>']
html+=['</section>']
for day in trip['days']:
 html+=['<section class="day"><p class="eyebrow">'+h(day['date']+' · '+day['city'])+'</p><h1>'+h(day['title'])+'</h1><img class="route" src="assets/'+posters[day['date']].name+'">']
 for kind,obj in cards(day):
  if kind=='optional':
   html+=['<article class="optional"><p class="small">有空再去 · 無固定時間</p><h3>'+h(obj.get('displayTitle',obj['title']))+'</h3><p>'+h(obj.get('description',''))+'</p>'+li(obj.get('mustTry',[]))+'<p>'+h(obj.get('note',''))+'</p>'+a(obj.get('mapUrl'),'位置／導航')+'</article>'];continue
  dining=obj['category'] in ['早餐','午餐','晚餐','甜點','條件甜點','餐飲散步'];lodging=bool(obj.get('lodgingAction'))
  title=hotel_name(obj,day);html+=['<article class="event '+('lodging' if lodging else 'dining' if dining else '')+'"><div class="time">'+h(zone_time(obj))+'<p class="small">'+h('住宿' if lodging else obj['category'])+'</p></div><div class="body"><h3>'+h(title)+'</h3>']
  if obj.get('status'):html+=['<span class="status '+('warning' if obj.get('statusKind')=='warning' else '')+'">'+h(obj['status'])+'</span>']
  rich=dining or obj.get('venue') or obj.get('isFlight')
  if rich:html+=[info(obj)]
  else:html+=['<p class="small">'+h(obj.get('activity',{}).get('description',''))+'</p>']
  for option in obj.get('mealOptions',[]):html+=['<section class="meal-option"><h3>'+h(option['title'])+'</h3><p>'+h(option['displayTime'])+'</p>'+info(option)+a(option.get('mapUrl'),'這間餐廳導航')+'</section>']
  html+=['<p>'+a(obj.get('mapUrl'),'導航')+(' · '+a(obj['ticketPlan']['url'],'官方／購票') if obj.get('ticketPlan') else '')+'</p></div></article>']
 html+=['</section>']
html+=['<section class="section"><h1>票券與訂位索引</h1><p>餐廳已訂位係使用者確認；沒有確認信者不生成可出示原票。地下教堂舊日期不能視為有效。所有訂位與原票請依票面核對。</p>']
for t in ledger:html+=['<article class="box"><h3>'+h(t['id']+' · '+t['title'])+'</h3><p>'+h(str(t['date'])+'｜'+t['time']+'｜'+t['status'])+'</p><p>'+h(t['note'])+'</p>'+a(t['url'],'原票／官方入口')+'</article>']
html+=['</section><section class="section"><h1>資料與費用說明</h1><p>本次由原Excel與最新App資料重編，完整更新時間軸、餐飲、導航與待訂清單。原Excel檔未覆寫；票券附件沿用既有同行包及雲端資料夾。</p><p>原1013預算頁保留為歷史參考，尚未完整重估新增博物館與餐廳費用。高鐵票面€121.20與使用者先前提供€100.75分列待核，不把票面差額直接改成已付。</p><p>畢卡索與奎爾宮營業資訊於10/22讀取官網；2027/1/2實際開放、價格與可售時段仍以供應商最新公告和最後票面為準。</p></section></body></html>']
handbook_html=P/f'Spain_2026_Travel_Handbook_{STAMP}.html';handbook_html.write_text(''.join(html))
handbook_pdf=P/f'Spain_2026_Travel_Handbook_{STAMP}.pdf'
class Handler(SimpleHTTPRequestHandler):
 def __init__(self,*a,**kw):super().__init__(*a,directory=str(P),**kw)
 def log_message(self,*_a):pass
server=ThreadingHTTPServer(('127.0.0.1',0),Handler);Thread(target=server.serve_forever,daemon=True).start()
try:
 with sync_playwright() as pw:
  browser=pw.chromium.launch(executable_path='/usr/bin/chromium',args=['--no-sandbox']);page=browser.new_page()
  page.goto(f'http://127.0.0.1:{server.server_port}/'+handbook_html.name,wait_until='networkidle');page.evaluate('document.fonts.ready.then(()=>true)')
  page.pdf(path=str(handbook_pdf),format='A4',print_background=True,display_header_footer=True,header_template='<div></div>',footer_template='<div style="width:100%;text-align:center;font-size:8px;color:#7a857b;">西班牙口袋旅程 · '+version+' · <span class="pageNumber"></span> / <span class="totalPages"></span></div>',margin={'top':'16mm','bottom':'18mm','left':'15mm','right':'15mm'})
  browser.close()
finally:server.shutdown();server.server_close()
# Verify structure and include a practical download bundle.
check=openpyxl.load_workbook(workbook,read_only=False,keep_vba=True)
assert len(check['每日地圖']._images)==len(trip['days'])
assert len(check['詳細時間流'].values.__iter__().__next__())>=17
assert any('Sagardi' in str(row[6]) for row in check['詳細時間流'].iter_rows(min_row=6,values_only=True))
assert all(check['詳細時間流'].cell(i,3).value is None and check['詳細時間流'].cell(i,4).value is None for i in range(6,check['詳細時間流'].max_row+1) if check['詳細時間流'].cell(i,6).value=='彈性備選')
with zipfile.ZipFile(workbook) as z:assert z.testzip() is None
outputs=[workbook,handbook_pdf,handbook_html,P/f'Spain_2026_Review_{STAMP}.md',P/f'Spain_2026_Review_{STAMP}.json',P/f'Spain_2026_Ticket_Links_{STAMP}.html']
bundle=P/f'Spain_2026_Updated_Documents_{STAMP}.zip'
with zipfile.ZipFile(bundle,'w',zipfile.ZIP_DEFLATED) as z:
 for file in outputs+list(posters.values()):z.write(file,str(file.relative_to(P)))
with zipfile.ZipFile(bundle) as z:assert z.testzip() is None
report={'version':version,'source':str(args.source),'source_sha256':hashlib.sha256(args.source.read_bytes()).hexdigest(),'workbook':str(workbook),'handbook':str(handbook_pdf),'bundle':str(bundle),'sheets':check.sheetnames,'timeline_events':sum(len(d['events']) for d in trip['days']),'optional_stops':sum(len(d.get('optionalStops',[])) for d in trip['days']),'route_diagrams':len(posters),'files':[{ 'name':f.name,'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in outputs+[bundle]]}
(P/'review/document-export-1022.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
