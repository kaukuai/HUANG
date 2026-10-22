"""Rebuild a private companion pack from a reviewed workbook and local originals.

Never writes into the repository or public site. No downloads, credentials or
private source IDs are embedded in this script.
"""
from pathlib import Path
from copy import deepcopy
from datetime import datetime, timezone
import argparse,base64,hashlib,json,mimetypes,re
import openpyxl
from travel_audit import review

ROOT=Path(__file__).resolve().parents[1]

def href(cell):
    if cell.hyperlink:return cell.hyperlink.target
    match=re.search(r'HYPERLINK\("([^"]+)',str(cell.value or ''))
    return match[1] if match else ''

def build(workbook,sources,output,folder_url='',international=None,vueling=None):
    output=output.resolve()
    if output.is_relative_to(ROOT):raise ValueError('私人資料包不得輸出到Git或公開網站內。')
    output.mkdir(parents=True,exist_ok=True)
    trip=json.loads((ROOT/'app/data/share.json').read_text())
    wb=openpyxl.load_workbook(workbook,read_only=False,data_only=False)
    source_files=json.loads(sources.read_text())
    if any(not source.get('downloaded') for source in source_files):raise ValueError('票券來源尚有下載失敗，請先補齊。')
    assets=[];by_source={}
    def asset(path,name=None,mime=None):
        path=Path(path);content=path.read_bytes();key=hashlib.sha256(content).hexdigest()[:24]
        if not any(a['id']==key for a in assets):
            mime=mime or ('application/vnd.apple.pkpass' if path.suffix=='.pkpass' else mimetypes.guess_type(path.name)[0])
            if mime not in ['application/pdf','image/png','image/jpeg','application/vnd.apple.pkpass','text/calendar']:raise ValueError('不支援的票券格式')
            assets.append({'id':key,'name':name or path.name,'type':mime,'data':base64.b64encode(content).decode()})
        return key
    for source in source_files:
        if source.get('downloaded'):
            if re.search(r'passport|護照',source['name'],re.I):raise ValueError('不可將證件加入票券包')
            path=Path(source['localPath']);content=path.read_bytes()
            if hashlib.sha256(content).hexdigest()!=source['sha256']:raise ValueError('原件校驗不符')
            by_source[source['id']]=asset(path)
    def folder_files(fragment):
        return [{'label':s['name'],'assetId':by_source.get(s['id']),'url':'https://drive.google.com/file/d/'+s['id']+'/view'} for s in source_files if fragment in s['path'] and s.get('downloaded')]
    groups={'K01':'馬德里皇宮/','K02':'塞維亞王宮/','K03':'塞維亞王宮/','K04':'塞維亞主教堂/','K05':'佛朗明哥表演/','K06':'阿爾布拉罕宮/','K07':'聖家堂和誕生面塔/','K08':'1228 格拉納達','K11':'巴特婁之家/','K12':'聖保祿醫院/','K13':'奎爾公園/','K14':'文森之家/','K15':'米拉之家/','K16':'奎爾紡織村＆地下教堂/','K17':'1226馬德里'}
    tickets=[]
    for row in wb['票券核對'].iter_rows(min_row=6):
        v=[c.value for c in row]
        if not v[0]:continue
        key=str(v[0]);status=str(v[8] or '')
        warning=key in ['K10','K16','P01','P02','P03'] or any(s in status for s in ['尚未','待改','未安排'])
        ticket={'id':key,'date':str(v[1] or ''),'title':str(v[2] or ''),'time':str(v[3] or ''),'report':str(v[4] or ''),'status':status,'statusKind':'warning' if warning else 'confirmed','note':str(v[9] or ''),'url':href(row[10]),'files':folder_files(groups[key]) if key in groups else []}
        if key=='K03':ticket['files'].sort(key=lambda f:'上層' not in f['label'])
        if key=='K02':ticket['files'].sort(key=lambda f:'一般票' not in f['label'])
        if key=='K17':ticket.update(status='已購票 · 原票已核對',price={'user_reported_total_eur':100.75,'face_total_eur':121.20,'price_difference_eur':20.45})
        if key=='K16':ticket.update(status='票券日期待改訂',note='目前原票仍印02/11/2026 10:00，非2027/1/2有效票。使用者已確認訂錯，正確新票待補；不推定含FGC。')
        if key=='F01':
            ticket['files']=[{'label':'四人國際機票行程單','assetId':asset(international),'url':''}] if international else []
        if key=='K09':
            ticket['files']=[{'label':'VY1121四人原行程單','assetId':asset(vueling),'url':''}] if vueling else []
            if not vueling:ticket.update(status='已購 · 本包原始附件待補',statusKind='warning',missingAttachment=True,note='VY1121已購，12/29 10:55–12:20。此批雲端來源未包含原行程單；新版App匯入時會保留此裝置舊包已有的VY1121附件。首次匯入而無舊原件的裝置仍需補原始PDF；不以自製文件充當原票。')
        tickets.append(ticket)
    ticket_by_id={t['id']:t for t in tickets}
    venue_ids={'Palacio Real de Madrid':['K01'],'Real Alcázar de Sevilla':['K02','K03'],'Catedral de Sevilla / Giralda':['K04'],'Iglesia del Salvador':['K04'],'Casa de la Memoria':['K05'],'Basílica de la Sagrada Família':['K07'],'Casa Batlló':['K11'],'Recinte Modernista de Sant Pau':['K12'],'Park Güell':['K13'],'Casa Vicens':['K14'],'Casa Milà / La Pedrera':['K15'],'Cripta de la Colònia Güell':['K16'],'Los Galayos':['RLOS'],'Botín':['RBOT'],'El Rinconcillo':['RELR'],'Bar Postiguillo Tapas':['RPOST']}
    direct={'5ec8ea0c04e4':['K17'],'10c64010c536':['K08'],'1927bf1e5d54':['K09']}
    for day in trip['days']:
        map_path=output.parent/'assets'/f'{day["date"]}-route.png'
        if not map_path.exists():raise FileNotFoundError('缺少最新每日動線圖：'+str(map_path))
        day['mapAsset']=asset(map_path,name=day['date']+'-每日動線示意.png')
        day['offlineMapNeedsUpdate']=False
        for event in day['events']:
            records=([event['reservation']] if event.get('reservation') else [])+[o['reservation'] for o in event.get('mealOptions',[]) if o.get('reservation')]+([event['ticketPlan']] if event.get('ticketPlan') else [])
            for record in records:
                if record['id'] not in ticket_by_id:
                    copied={**deepcopy(record),'files':[]};tickets.append(copied);ticket_by_id[copied['id']]=copied
            name=event.get('venue',{}).get('nameEn','')
            ids=[r['id'] for r in records] or direct.get(event['id']) or (['K06'] if name.startswith('Alhambra /') else venue_ids.get(name,[]))
            if event.get('isFlight'):ids=['F01']
            event['ticketIds']=ids
            linked=[ticket_by_id[i] for i in ids if i in ticket_by_id]
            if linked and not event.get('status'):
                event.update(status=linked[0]['status'],statusKind=linked[0]['statusKind'])
    trip['hotels']=[]
    for row in wb['交通與住宿'].iter_rows(min_row=6,max_row=9):
        v=[c.value for c in row]
        trip['hotels'].append({'city':str(v[0]),'name':str(v[1]),'dates':str(v[2]),'address':str(v[4]),'note':str(v[6]),'mapUrl':href(row[8])})
    for day in trip['days']:
        matches=[h for h in trip['hotels'] if h['city'] in day['city']]
        day['hotel']=matches[-1]['name'] if matches else ''
    audit=review(trip)
    trip['todos']=[{'title':r['item'],'note':r['date']+'；'+r['status']+'。'+r['action'],'url':r['url']} for r in audit['attractions']+audit['restaurants_pending']+audit['transport'][4:]]
    for r in trip.get('planningReminders',[]):
        if not r.get('calendar'):continue
        # Use the same tested calendar generator as the app below (calendar file written by caller).
        path=output/(r['id']+'.ics')
        if path.exists():asset(path,mime='text/calendar')
    trip.update(tickets=sorted(tickets,key=lambda t:(t['date'] or '9999',t['id'])),flights=[],budget={'core':None,'allIn':None,'note':'最新餐飲、門票與交通費尚未重新對帳，不使用原Excel的舊預算冒充最新總額。'},ticketReview={'status':'current','last_verified':datetime.now(timezone.utc).date().isoformat(),'folderUrl':folder_url,'note':'雲端28份票券／附件已重新下載；VY1121原件不在此批來源，詳見K09；地下教堂日期錯誤仍待改訂。'})
    pack={'format':'spain-pocket-trip/v1','createdAt':datetime.now(timezone.utc).isoformat(),'retainExistingTicketIds':['K09','F01'],'trip':trip,'assets':assets}
    path=output/f'Spain_Travel_{trip["version"]}.travelpack.json'
    path.write_text(json.dumps(pack,ensure_ascii=False,separators=(',',':')))
    report={'private_pack':str(path),'planning_version':trip['version'],'days':len(trip['days']),'events':sum(len(d['events']) for d in trip['days']),'tickets':len(tickets),'evidence_files':len(by_source)+(1 if international else 0)+(1 if vueling else 0),'maps':len(trip['days']),'assets':len(assets),'private_bytes':path.stat().st_size,'ticket_review_status':'current','missing_originals':[] if vueling else ['K09 VY1121原始PDF'],'retained_device_ticket_ids':pack['retainExistingTicketIds']}
    (output/'pack-manifest.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    (output.parent/'review/app-build-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    return report

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--workbook',type=Path,required=True);parser.add_argument('--sources',type=Path,required=True);parser.add_argument('--output',type=Path,required=True);parser.add_argument('--drive-folder-url',default='');parser.add_argument('--international',type=Path);parser.add_argument('--vueling',type=Path)
    a=parser.parse_args();print(json.dumps(build(a.workbook,a.sources,a.output,a.drive_folder_url,a.international,a.vueling),ensure_ascii=False,indent=2))
