"""Build a public itinerary app and a separate, device-imported private travel pack."""
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from urllib.parse import unquote, quote_plus
import argparse, base64, hashlib, json, mimetypes, posixpath, re, shutil
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--planning', type=Path, default=Path('/workspace/spain-planning'))
parser.add_argument('--output', type=Path, default=ROOT/'dist')
args = parser.parse_args()
P, OUT = args.planning, args.output
state = json.loads((P/'latest-planning-state.json').read_text())
master = json.loads((P/state['master']).read_text())
workbook = P/state['timeline']
values = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
formulas = openpyxl.load_workbook(workbook, read_only=True, data_only=False)
ZONE = ZoneInfo('Europe/Madrid')
assets, by_path, by_name = [], {}, {}
def asset(path=None, content=None, name=None, mime=None):
    if path is not None:
        path=Path(path);content=path.read_bytes();name=name or path.name
    key=hashlib.sha256(content).hexdigest()[:24]
    if not any(a['id']==key for a in assets):
        mime=mime or mimetypes.guess_type(name)[0] or 'application/octet-stream'
        if str(name).endswith('.pkpass'):mime='application/vnd.apple.pkpass'
        assets.append({'id':key,'name':name,'type':mime,'data':base64.b64encode(content).decode()})
    if path is not None:
        by_path[str(path.resolve())]=key;by_name.setdefault(path.name,key)
    return key
evidence=sorted(f for f in (P/'booking-evidence/2026-10-14').rglob('*') if f.is_file())
for path in evidence:
    assert not re.search(r'passport|護照',path.name,re.I)
    asset(path)
def href(cell):
    text=str(cell.value or '')
    m=re.search(r'HYPERLINK\("((?:[^"]|"")+)"',text)
    return m.group(1).replace('""','"') if m else (cell.hyperlink.target if getattr(cell,'hyperlink',None) else '')
def txt(value):return str(value or '').strip()
def maps(query):return 'https://www.google.com/maps/search/?api=1&query='+quote_plus(query)

days=[]
for row in values['行程總覽'].iter_rows(min_row=6,max_row=16,values_only=True):
    date=row[0].strftime('%Y-%m-%d')
    days.append({'date':date,'city':txt(row[2]),'title':txt(row[3]),'hotel':txt(row[7]),'events':[],'mapAsset':None})
day_by_date={d['date']:d for d in days}
for row, fr in zip(values['詳細時間流'].iter_rows(min_row=6), formulas['詳細時間流'].iter_rows(min_row=6)):
    v=[c.value for c in row]
    if len(v)<7 or not isinstance(v[0],datetime) or not isinstance(v[2],datetime):continue
    date=v[0].strftime('%Y-%m-%d');title=txt(v[6]);note=txt(v[8]);category=txt(v[5])
    conditional=any(s in title+note for s in ['有餘裕','開門才去','有開才去','條件行程','條件甜點','彈性安排'])
    ids=re.findall(r'K\d{2}',txt(v[9]))
    status,kind='','neutral'
    if conditional:status,kind='彈性／條件行程','conditional'
    elif '未購' in txt(v[12]) or (category not in ['交通','休息'] and ('尚未購票' in note or '目標尚未購' in note)):status,kind='尚未購票','warning'
    elif any(s in note for s in ['尚未訂','未訂位','晚餐尚未安排','目前未聯絡到主人']):status,kind='尚未安排','warning'
    elif '待使用者改訂' in note:status,kind='票券日期待改訂','warning'
    elif category not in ['交通','休息'] and (ids or any(s in title for s in ['Los Galayos','Botín','El Rinconcillo','Madrid→Sevilla'])):
        status,kind=('已確認','confirmed')
    event={'id':hashlib.sha1(f'{date}|{v[2].isoformat()}|{title}'.encode()).hexdigest()[:12],
           'date':date,'start':v[2].replace(tzinfo=ZONE).isoformat(),'end':v[3].replace(tzinfo=ZONE).isoformat(),
           'category':category,'title':title,'transport':txt(v[7]),'note':note,'ticketIds':sorted(set(ids)),
           'mapUrl':href(fr[10]),'infoUrl':href(fr[15]),'conditional':conditional,'status':status,'statusKind':kind}
    if 'Prado名畫' in title:event.update(status='尚未購票',statusKind='warning')
    if 'Murillo' in title and category=='午餐':event.update(status='尚未訂位',statusKind='warning')
    if 'Madrid→Sevilla'==title:event['ticketIds']=['K17']
    if '聖米格爾市場' in title and category=='午餐':event.update(status='聖誕營業待確認',statusKind='warning')
    day_by_date[date]['events'].append(event)
assert sum(len(d['events']) for d in days)==211

# Read each drawing relationship instead of assuming image numbering.
with ZipFile(workbook) as z:
    n={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships','p':'http://schemas.openxmlformats.org/package/2006/relationships','x':'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    root=ET.fromstring(z.read('xl/worksheets/sheet13.xml'))
    rid=root.find('s:drawing',n).attrib['{'+n['r']+'}id']
    rels=ET.fromstring(z.read('xl/worksheets/_rels/sheet13.xml.rels'))
    target=next(r.attrib['Target'] for r in rels if r.attrib['Id']==rid)
    drawing=posixpath.normpath(posixpath.join('xl/worksheets',target))
    relpath=posixpath.join(posixpath.dirname(drawing),'_rels',posixpath.basename(drawing)+'.rels')
    drawing_rels={r.attrib['Id']:posixpath.normpath(posixpath.join(posixpath.dirname(drawing),r.attrib['Target'])) for r in ET.fromstring(z.read(relpath))}
    anchors=[]
    for a in ET.fromstring(z.read(drawing)):
        row=a.find('x:from/x:row',n);blip=a.find('.//a:blip',n)
        if row is not None and blip is not None:anchors.append((int(row.text),drawing_rels[blip.attrib['{'+n['r']+'}embed']]))
    assert len(anchors)==11
    for day,(_,member) in zip(days,sorted(anchors)):
        suffix=Path(member).suffix
        day['mapAsset']=asset(content=z.read(member),name=f'{day["date"]}-每日路線圖{suffix}')

drive_files=[]
for f in master['drive_ticket_links']:
    candidate=by_name.get(f['name']) or by_name.get(f['name']+'.xlsx')
    drive_files.append({'label':f['name'],'url':f['url'],'assetId':candidate,'folder':str(Path(f['path']).parent)})
drive_by_url={f['url']:f for f in drive_files}
def local_file(name,label):return {'label':label,'assetId':by_name[name],'url':''}
tickets=[]
for row,fr in zip(values['票券核對'].iter_rows(min_row=6),formulas['票券核對'].iter_rows(min_row=6)):
    v=[c.value for c in row]
    if not v or not re.fullmatch(r'K\d{2}',txt(v[0])):continue
    url=href(fr[10]);files=[]
    if url in drive_by_url:
        selected=drive_by_url[url]
        files=[{k:f[k] for k in ['label','url','assetId']} for f in drive_files if f['folder']==selected['folder']]
        # Alcázar's two distinct tickets keep their primary file first.
        files.sort(key=lambda f:f['url']!=url)
    if v[0]=='K08':files=[local_file('1228車票.pdf','ALSA四人巴士票')]
    if v[0]=='K09':files=[local_file('電子行程單 (已含英文版本)2.pdf','VY1121四人行程單'),local_file('電子收據 (已含英文版本)2.pdf','VY1121付款收據')]
    pending=v[0] in ['K10','K16']
    tickets.append({'id':v[0],'date':txt(v[1]),'title':txt(v[2]),'time':txt(v[3]),'report':txt(v[4]),
                    'status':('票券日期待改訂' if v[0]=='K16' else '晚餐尚未安排' if v[0]=='K10' else '已購 · 原票待補' if v[0]=='K17' else '已購票'),
                    'statusKind':'warning' if pending else 'confirmed','note':txt(v[9]),'url':url if url.startswith('https://') else '', 'files':files})
tickets.append({'id':'F01','date':'2026-12-23','title':'國際機票 · 中國東方航空','time':'12/23出發 · 1/4返抵台灣','report':'各航段依機場當地時間',
                'status':'已開票','statusKind':'confirmed','note':'四人同一訂單，國際航段每人一件23kg托運、一件8kg手提及一件座下個人物品。請依航司最新通知確認航廈及截止時間。',
                'files':[local_file('電子行程單.pdf','四人國際機票行程單'),local_file('電子收據 (已含英文版本).pdf','國際機票付款收據')],'url':''})
for i,title,date,time,url in [('P01','普拉多美術館','2026-12-26','15:45–17:00目標','https://entradas.museodelprado.es/'),('P02','奎爾宮 Palau Güell','2027-01-02','14:30目標','https://inici.palauguell.cat/en/')]:
    tickets.append({'id':i,'title':title,'date':date,'time':time,'report':'待購票確認','status':'尚未購票','statusKind':'warning','note':'目前是規劃時段，尚未取得有效入場票。','files':[],'url':url})
tickets.sort(key=lambda t:(t['date'],t['id']))
for tid,title,date,time,note in [
    ('RLOS','Los Galayos晚餐','2026-12-25','19:00','使用者確認已訂位；18:50到店目標，晚餐前順逛Plaza Mayor。確認信尚未提供。'),
    ('RBOT','Botín午餐','2026-12-26','13:00','使用者確認已訂位，預計15:00用餐結束，再前往Prado。確認信尚未提供。'),
    ('RELR','El Rinconcillo晚餐','2026-12-26','22:30','使用者確認已訂位，只保留十分鐘至22:40。21:39列車抵達後先入住，銜接偏緊，需預辦入住與接車。確認信尚未提供。'),
]:
    tickets.append({'id':tid,'title':title,'date':date,'time':time,'report':'依行程提早到店','status':'已訂位 · 使用者確認','statusKind':'confirmed','note':note,'files':[],'url':''})
for d in days:
    for e in d['events']:
        for name,tid in [('Los Galayos','RLOS'),('Botín','RBOT'),('El Rinconcillo','RELR')]:
            if e['title']==name:e['ticketIds']=[tid]
tickets.append({'id':'REF','title':'購票清單與對照表','date':'','time':'票券資料參考','report':'','status':'參考資料','statusKind':'neutral','note':'清單用於對照，不是景點入場票。','files':[{k:f[k] for k in ['label','url','assetId']} for f in drive_files if f['label']=='SPAIN TICKETS LIST'],'url':''})
tickets.sort(key=lambda t:(t['date'] or '9999',t['id']))
flights=[]
for f in master['air']['segments']:
    if f['flight_number']=='VY1121':continue
    dep,arr=f['departure'],f['arrival']
    start=datetime.fromisoformat(dep['date_time_local']).replace(tzinfo=ZoneInfo(dep['timezone'])).isoformat()
    end=datetime.fromisoformat(arr['date_time_local']).replace(tzinfo=ZoneInfo(arr['timezone'])).isoformat()
    flights.append({'id':'flight-'+f['flight_number'],'number':f['flight_number'],'title':f'{f["flight_number"]} · {dep["airport_iata"]} → {arr["airport_iata"]}',
                    'start':start,'end':end,'date':dep['date_time_local'][:10],'category':'航班','isFlight':True,
                    'displayTime':f'{dep["airport_iata"]} {dep["date_time_local"][11:16]} → {arr["airport_iata"]} {arr["date_time_local"][11:16]}',
                    'status':'已開票','statusKind':'confirmed','conditional':False,'ticketIds':['F01'],'mapUrl':maps(dep['airport_iata']+' Airport '+str(dep['terminal'] or '')),
                    'note':f'顯示時間各為機場當地時間。出發：{dep["date_time_local"][:10]} {dep["terminal"] or "航廈待確認"}；抵達：{arr["date_time_local"][:10]} {arr["terminal"] or "航廈待確認"}。'})
hotels=[]
for row,fr in zip(values['交通與住宿'].iter_rows(min_row=6,max_row=9),formulas['交通與住宿'].iter_rows(min_row=6,max_row=9)):
    v=[c.value for c in row]
    hotels.append({'city':txt(v[0]),'name':txt(v[1]),'dates':txt(v[2]),'address':txt(v[4]),'note':txt(v[6]),'mapUrl':href(fr[8])})
todos=[
 {'title':'11/25開始預約Murillo午餐','note':'12/25四人13:00目標；僅提前一個月開放為使用者提供，精確放位時間待確認。','url':'https://murillocafe.com/'},
 {'title':'Prado與奎爾宮仍待購票','note':'Prado12/26 15:45目標；奎爾宮1/2 14:30目標。','url':'https://entradas.museodelprado.es/'},
 {'title':'地下教堂票券改訂日期','note':'原票日期訂錯，請改成2027/1/2並補新票；目前不能把舊票當作1/2有效票。','url':''},
 {'title':'塞維亞先入住、再赴22:30晚餐','note':'21:39抵達後只有51分鐘；預辦住宿登記與門碼、安排接車。22:40僅為訂位保留截止。','url':''},
 {'title':'補上高鐵原票','note':'19:00–21:39已購，四人總€100.75；車次、業者與實際站名待票面確認。','url':''},
 {'title':'12/24晚餐與聖誕營業','note':'主人尚未聯絡到，未成行在飯店自理；市場、超市部與早到寄物仍需確認。','url':''},
]
trip={'schema':1,'version':master['version'],'title':'西班牙・口袋旅程','timezone':'Europe/Madrid','people':4,
      'days':days,'tickets':tickets,'hotels':hotels,'todos':todos,'flights':flights,
      'budget':{'core':master['budget']['core_including_domestic_excluding_international_per_person_twd'],
                'allIn':master['budget']['including_international_per_person_twd']}}

# Public mode contains no evidence, lodging addresses, booking IDs or private links.
share=json.loads(json.dumps(trip))
share.update(tickets=[],hotels=[],todos=[],flights=[],budget=None)
hotel_pattern=re.compile(r'Arenal(?: Suites(?: Puerta del Sol)?)?|Enjoy\s*Sevilla|EnjoySevilla|Sercotel(?: Granada Suites)?|Inside Barcelona Apartments Príncep|Príncep',re.I)
for day in share['days']:
    day['hotel']='';day['mapAsset']=None
    for e in day['events']:
        sensitive=bool(hotel_pattern.search(e['title']+' '+unquote(e['mapUrl']))) or any(s in e['title'] for s in ['住宿','飯店','公寓','主人'])
        e['title']=hotel_pattern.sub('住宿',e['title'])
        e.update(note='',transport='',ticketIds=[],infoUrl='')
        if sensitive:e['mapUrl']=''
        if e['statusKind']=='confirmed':e.update(status='',statusKind='neutral')
        if '主人' in e['title']:e['title']='自由安排／住宿休息'
        if any(word in e['title'] for word in ['護照','passport']):e['title']='行前準備／整理用品'
share_json=json.dumps(share,ensure_ascii=False,separators=(',',':'))
for forbidden in ['drive.google.com','booking-evidence','Calle del Arenal','Plaza Aviador','passport','護照']:
    assert forbidden not in share_json,forbidden
(ROOT/'app/data').mkdir(exist_ok=True)
(ROOT/'app/data/share.json').write_text(json.dumps(share,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
OUT.mkdir(parents=True,exist_ok=True)
for path in (ROOT/'app').iterdir():
    if path.is_file():shutil.copy2(path,OUT/path.name)
(OUT/'data').mkdir(exist_ok=True)
(OUT/'data/share.json').write_text(share_json,encoding='utf-8')
build_hash=hashlib.sha256(share_json.encode())
for source in sorted((ROOT/'app').iterdir()):
    if source.is_file():build_hash.update(source.read_bytes())
build_id=build_hash.hexdigest()[:12]
(OUT/'sw.js').write_text((ROOT/'app/sw.js').read_text().replace('__BUILD_ID__',build_id))
private_dir=P/'app-private';private_dir.mkdir(exist_ok=True)
pack_path=private_dir/f'Spain_Travel_{master["version"]}.travelpack.json'
pack_path.write_text(json.dumps({'format':'spain-pocket-trip/v1','trip':trip,'assets':assets},ensure_ascii=False,separators=(',',':')),encoding='utf-8')
# Icon rendering from the source SVG uses Chromium in the browser test/build finalization.
report={'public_dir':str(OUT),'private_pack':str(pack_path),'days':len(days),'events':sum(len(d['events']) for d in days),
        'tickets':len(tickets),'evidence_files':len(evidence),'maps':len(anchors),'assets':len(assets),'private_bytes':pack_path.stat().st_size,'build_id':build_id}
(P/'review/app-build-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
